import openpyxl

wb = openpyxl.load_workbook('Team Details (1).xlsx')
for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    for r_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        for c_idx, val in enumerate(row, start=1):
            if val is not None:
                val_str = str(val).lower()
                for name in ['fardeen', 'dheeraj', 'sajja', 'nagaraju', 'pravallika', 'ayushi']:
                    if name in val_str:
                        print(f"Sheet: {sheet_name}, Cell: {r_idx}:{c_idx}, Val: {val}")
