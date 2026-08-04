import openpyxl

wb = openpyxl.load_workbook('Team Details (1).xlsx')
for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        continue
    # find headers
    header_row = None
    for r in rows[:5]:
        if any(isinstance(x, str) and ('email' in x.lower() or 'mail' in x.lower()) for x in r if x is not None):
            header_row = r
            break
    if header_row:
        print(f"Sheet: {sheet_name} has email columns. Headers: {header_row[:12]}")
