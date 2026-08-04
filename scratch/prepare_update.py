import openpyxl

wb = openpyxl.load_workbook('Team Details (1).xlsx')

whitelist = ['fardeen', 'dheeraj', 'karthik sajja', 'karthik', 'sajja', 'nagaraju', 'pravallika', 'ayushi']

def is_whitelisted(name_val):
    if name_val is None:
        return False
    name_str = str(name_val).lower()
    return any(w in name_str for w in whitelist)

for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        continue
    
    # Identify headers and name/email columns
    email_cols = []
    name_col = None
    
    # Look for name and email headers in first few rows
    for r_idx, r in enumerate(rows[:5]):
        has_email = False
        for c_idx, val in enumerate(r):
            if val is not None:
                val_str = str(val).lower()
                if 'email' in val_str or 'mail' in val_str:
                    email_cols.append(c_idx)
                    has_email = True
                if 'name' in val_str and name_col is None:
                    name_col = c_idx
        if has_email:
            break
            
    if not email_cols:
        continue
        
    print(f"\n--- Sheet: {sheet_name} (Name Col: {name_col}, Email Cols: {email_cols}) ---")
    for r_idx, r in enumerate(rows):
        # Check if row has any email
        has_any_email = any(r[c] is not None for c in email_cols)
        if not has_any_email:
            continue
            
        name_val = r[name_col] if name_col is not None else None
        keep = is_whitelisted(name_val)
        
        # print status
        emails = [r[c] for c in email_cols]
        print(f"Row {r_idx+1} | Name: {name_val} | Emails: {emails} | Keep: {keep}")
