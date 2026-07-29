import openpyxl

wb = openpyxl.load_workbook('Team Details (1).xlsx')
for name in ['Team Details', 'Team_contactDetails', 'Distribution']:
    if name in wb.sheetnames:
        sheet = wb[name]
        print(f"--- Sheet: {name} ---")
        rows = list(sheet.iter_rows(values_only=True))
        for idx, r in enumerate(rows):
            # Print row index and the row if it contains anything
            if any(r):
                print(f"Row {idx}: {r[:10]}")
